"""Unit and integration tests for Phase 12B:
- Deterministic resume linter (all errors, warnings, and tips)
- ATS-safe Jinja2 templates (classic, compact, technical) with HTML injection safety
- Multi-format rendering (HTML, PDF, DOCX, Markdown, Text, JSON)
- PDF ATS readability verification via pdfplumber (>95% skill recovery)
- DOCX structure and heading verification
- Deterministic resume import from PDF, DOCX, and TXT
- FastAPI /resume endpoints (catalog, socials, lint, render, import)
"""

from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Any, Dict, List

import docx
import pdfplumber
import pytest
from fastapi.testclient import TestClient

from api.main import app
from ml.resume_import import import_resume
from ml.resume_lint import lint
from ml.resume_render import (
    export_json,
    group_skills,
    prepare_links,
    render_docx,
    render_html,
    render_markdown,
    render_pdf,
    render_text,
)
from ml.resume_schema import Resume

FIXTURE_1 = Path(__file__).resolve().parent / "fixtures" / "resume_example_1.json"
FIXTURE_2 = Path(__file__).resolve().parent / "fixtures" / "resume_example_2.json"

client = TestClient(app)


@pytest.fixture
def resume_data_1() -> Dict[str, Any]:
    with open(FIXTURE_1, "r", encoding="utf-8") as f:
        return json.load(f)


@pytest.fixture
def resume_data_2() -> Dict[str, Any]:
    with open(FIXTURE_2, "r", encoding="utf-8") as f:
        return json.load(f)


# =========================================================================
# 1. Resume Linter Tests
# =========================================================================

def test_lint_detects_errors(resume_data_1):
    """Test critical error detection in resume linter."""
    bad_resume = json.loads(json.dumps(resume_data_1))
    bad_resume["basics"]["name"] = ""
    bad_resume["basics"]["email"] = "invalid-email"
    bad_resume["basics"]["phone"] = "abc12"
    bad_resume["experience"][0]["start"] = "2023-01"
    bad_resume["experience"][0]["end"] = "2021-01"
    bad_resume["experience"][0]["company"] = ""

    issues = lint(bad_resume)
    errors = [i for i in issues if i["level"] == "error"]
    err_paths = {i["path"] for i in errors}

    assert "basics.name" in err_paths
    assert "basics.email" in err_paths
    assert "basics.phone" in err_paths
    assert "experience[0].end" in err_paths
    assert "experience[0].company" in err_paths


def test_lint_detects_bullet_length_error(resume_data_1):
    """Test bullet over 300 characters error."""
    bad_resume = json.loads(json.dumps(resume_data_1))
    bad_resume["experience"][0]["bullets"].append("A" * 350)

    issues = lint(bad_resume)
    errors = [i for i in issues if i["level"] == "error"]
    assert any("300 character limit" in e["message"] for e in errors)


def test_lint_detects_warnings(resume_data_1):
    """Test style and structural warning detection."""
    warn_resume = json.loads(json.dumps(resume_data_1))

    # Weak phrase
    warn_resume["experience"][0]["bullets"] = [
        "Responsible for writing backend microservices.",
        "I developed an API gateway for our services.",
    ]
    # Inconsistent dates
    warn_resume["experience"][0]["start"] = "2021-05"
    warn_resume["education"][0]["start"] = "2017"
    # Entry with no bullets
    warn_resume["experience"][1]["bullets"] = []
    # Duplicate skill
    warn_resume["skills"].append({
        "category_id": "languages",
        "name": "Python",
        "taxonomy_id": "python",
    })
    # Non-resume friendly link
    warn_resume["basics"]["links"].append({
        "platform_id": "discord",
        "value": "cooldev#1234",
    })

    issues = lint(warn_resume)
    warnings = [i for i in issues if i["level"] == "warning"]
    messages = " ".join(w["message"] for w in warnings)

    assert "weak phrase" in messages.lower()
    assert "first-person pronouns" in messages.lower()
    assert "inconsistent date formats" in messages.lower()
    assert "no bullet points" in messages.lower()
    assert "duplicate skill" in messages.lower()
    assert "non-resume-friendly" in messages.lower()


def test_lint_detects_tips(resume_data_1):
    """Test tips detection (metrics, summary length, repeated verbs)."""
    tip_resume = json.loads(json.dumps(resume_data_1))
    # Bullets without metrics
    tip_resume["experience"][0]["bullets"] = [
        "Engineered backend microservices and databases.",
        "Engineered caching layer for low latency.",
        "Engineered deployment pipelines with Docker.",
    ]
    # Summary with > 4 sentences
    tip_resume["summary"] = "First sentence. Second sentence. Third sentence. Fourth sentence. Fifth sentence."
    # Fewer than 5 skills
    tip_resume["skills"] = tip_resume["skills"][:3]

    issues = lint(tip_resume)
    tips = [i for i in issues if i["level"] == "tip"]
    messages = " ".join(t["message"] for t in tips)

    assert "quantifiable outcome" in messages.lower() or "number" in messages.lower()
    assert "sentences long" in messages.lower()
    assert "fewer" in messages.lower() or "listed 3 skills" in messages.lower()
    assert "used 3 times" in messages.lower()


# =========================================================================
# 2. Template Rendering & HTML Safety
# =========================================================================

@pytest.mark.parametrize("template_name", ["classic", "compact", "technical"])
@pytest.mark.parametrize("page_size", ["A4", "Letter"])
def test_render_html_templates(resume_data_1, template_name, page_size):
    """Assert all HTML templates render non-empty text containing key fields."""
    html_out = render_html(resume_data_1, template_name=template_name, page_size=page_size)
    assert len(html_out) > 500
    assert resume_data_1["basics"]["name"] in html_out
    assert resume_data_1["basics"]["email"] in html_out
    assert "Experience" in html_out
    assert "Education" in html_out
    assert "Skills" in html_out
    assert page_size in html_out


def test_html_injection_prevention(resume_data_1):
    """Verify that malicious script tags in input are safely escaped in HTML output."""
    injected = json.loads(json.dumps(resume_data_1))
    injected["basics"]["name"] = "<script>alert('xss')</script>Nikhil"
    injected["experience"][0]["company"] = "<img src=x onerror=alert(1)>TechCorp"

    html_out = render_html(injected, template_name="classic")
    assert "<script>" not in html_out
    assert "&lt;script&gt;" in html_out or "alert(&#39;xss&#39;)" in html_out
    assert "<img src=x" not in html_out


# =========================================================================
# 3. PDF Rendering & ATS Readability via pdfplumber
# =========================================================================

def test_render_pdf_and_verify_ats_recovery(resume_data_1):
    """Generate PDF and assert pdfplumber recovers >95% of candidate skills and core text."""
    pdf_bytes = render_pdf(resume_data_1, template_name="classic", page_size="A4")
    assert len(pdf_bytes) > 10000

    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        assert len(pdf.pages) >= 1
        extracted_text = "\n".join(page.extract_text() or "" for page in pdf.pages)

    # Verify candidate core info
    assert "NIKHIL MAGHODIYA" in extracted_text.upper()
    assert resume_data_1["basics"]["email"].lower() in extracted_text.lower()
    assert "EXPERIENCE" in extracted_text.upper()
    assert "EDUCATION" in extracted_text.upper()
    assert "SKILLS" in extracted_text.upper()

    # Verify skill recovery >= 95%
    candidate_skills = [s["name"] for s in resume_data_1["skills"]]
    recovered_count = sum(1 for s in candidate_skills if s.lower() in extracted_text.lower())
    recovery_rate = recovered_count / len(candidate_skills)
    assert recovery_rate >= 0.95, f"Skill recovery rate {recovery_rate:.1%} < 95% ({recovered_count}/{len(candidate_skills)})"


def test_render_pdf_fixture_2(resume_data_2):
    """Generate PDF for fixture 2 and verify ATS readability."""
    pdf_bytes = render_pdf(resume_data_2, template_name="compact", page_size="A4")
    assert len(pdf_bytes) > 10000

    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)

    assert resume_data_2["basics"]["name"].upper() in text.upper()
    assert resume_data_2["basics"]["email"].lower() in text.lower()


# =========================================================================
# 4. DOCX, Markdown, Text, and JSON Rendering
# =========================================================================

def test_render_docx_structure(resume_data_1):
    """Verify DOCX generates with correct Word headings and structure."""
    docx_bytes = render_docx(resume_data_1)
    assert len(docx_bytes) > 5000

    doc = docx.Document(io.BytesIO(docx_bytes))
    headings_1 = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading 1")]
    headings_2 = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading 2")]

    assert resume_data_1["basics"]["name"] in headings_1
    assert "Experience" in headings_2
    assert "Education" in headings_2
    assert "Skills" in headings_2
    assert "Projects" in headings_2


def test_render_markdown_and_text(resume_data_1):
    """Verify Markdown and Plaintext generation."""
    md = render_markdown(resume_data_1)
    assert md.startswith(f"# {resume_data_1['basics']['name']}")
    assert "## Experience" in md
    assert "## Skills" in md

    txt = render_text(resume_data_1)
    assert resume_data_1["basics"]["name"].upper() in txt
    assert "EXPERIENCE" in txt
    assert "SKILLS" in txt
    assert "  *" in txt  # ASCII bullet points


def test_export_json_validates(resume_data_1):
    """Verify exported JSON is valid and roundtrips to Resume schema."""
    json_str = export_json(resume_data_1)
    data = json.loads(json_str)
    res = Resume.model_validate(data)
    assert res.basics.name == resume_data_1["basics"]["name"]


# =========================================================================
# 5. Deterministic Resume Import
# =========================================================================

def test_import_resume_from_pdf(resume_data_1):
    """Round-trip test: render PDF then import, verifying core fields extracted."""
    pdf_bytes = render_pdf(resume_data_1, template_name="classic")
    result = import_resume(pdf_bytes, "test_resume.pdf")

    assert "resume" in result
    assert "needs_review" in result
    draft = result["resume"]

    # Name and email recovered
    assert "Nikhil" in draft["basics"]["name"]
    assert draft["basics"]["email"] == resume_data_1["basics"]["email"]
    assert len(draft["skills"]) > 5


def test_import_resume_from_docx(resume_data_1):
    """Round-trip test: render DOCX then import, verifying core fields extracted."""
    docx_bytes = render_docx(resume_data_1)
    result = import_resume(docx_bytes, "test_resume.docx")

    draft = result["resume"]
    assert "Nikhil" in draft["basics"]["name"]
    assert draft["basics"]["email"] == resume_data_1["basics"]["email"]
    assert len(draft["skills"]) > 5


# =========================================================================
# 6. API Endpoints
# =========================================================================

def test_api_get_catalog():
    """Verify /resume/catalog endpoint returns 15 categories."""
    resp = client.get("/resume/catalog")
    assert resp.status_code == 200
    data = resp.json()
    assert "categories" in data
    assert len(data["categories"]) == 15


def test_api_get_socials():
    """Verify /resume/socials endpoint returns supported platforms."""
    resp = client.get("/resume/socials")
    assert resp.status_code == 200
    platforms = resp.json()
    assert isinstance(platforms, list)
    assert any(p["id"] == "github" for p in platforms)


def test_api_post_lint(resume_data_1):
    """Verify /resume/lint endpoint accepts resume and returns issues list."""
    resp = client.post("/resume/lint", json=resume_data_1)
    assert resp.status_code == 200
    issues = resp.json()
    assert isinstance(issues, list)


@pytest.mark.parametrize("fmt", ["pdf", "html", "docx", "txt", "md", "json"])
def test_api_post_render(resume_data_1, fmt):
    """Verify /resume/render endpoint generates all requested formats."""
    payload = {
        "resume": resume_data_1,
        "format": fmt,
        "template": "classic",
        "page_size": "A4",
    }
    resp = client.post("/resume/render", json=payload)
    assert resp.status_code == 200
    assert len(resp.content) > 100


def test_api_post_import(resume_data_1):
    """Verify /resume/import endpoint parses multipart file upload."""
    pdf_bytes = render_pdf(resume_data_1, template_name="classic")
    files = {"file": ("test.pdf", pdf_bytes, "application/pdf")}
    resp = client.post("/resume/import", files=files)
    assert resp.status_code == 200
    data = resp.json()
    assert "resume" in data
    assert "needs_review" in data
    assert data["resume"]["basics"]["email"] == resume_data_1["basics"]["email"]
