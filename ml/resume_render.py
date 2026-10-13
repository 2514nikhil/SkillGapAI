"""Resume Rendering Engine for SkillGap.

Produces deterministic, ATS-compliant outputs in multiple formats:
- HTML (Jinja2 templates: classic, compact, technical)
- PDF (WeasyPrint with headless browser fallback)
- DOCX (python-docx with real Word heading styles)
- Markdown (plain GitHub-flavored text)
- Text (UTF-8 formatted plaintext with ASCII bullets)
- JSON (schema-compliant standard JSON)
"""

from __future__ import annotations

import io
import json
import logging
import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from jinja2 import Environment, FileSystemLoader, select_autoescape

from ml.resume_schema import Resume
from ml.socials import build_url, get_platform, load_platforms

logger = logging.getLogger(__name__)

TEMPLATES_DIR = Path(__file__).resolve().parent / "templates"
CATALOG_PATH = Path(__file__).resolve().parent.parent / "data" / "processed" / "tech_catalog.json"

_jinja_env: Optional[Environment] = None
_tech_catalog_cache: Optional[Dict[str, Any]] = None


def _get_jinja_env() -> Environment:
    global _jinja_env
    if _jinja_env is None:
        _jinja_env = Environment(
            loader=FileSystemLoader(str(TEMPLATES_DIR)),
            autoescape=select_autoescape(["html", "xml"]),
            trim_blocks=True,
            lstrip_blocks=True,
        )
    return _jinja_env


def _load_catalog() -> Dict[str, Any]:
    global _tech_catalog_cache
    if _tech_catalog_cache is None:
        if CATALOG_PATH.exists():
            with open(CATALOG_PATH, "r", encoding="utf-8") as f:
                _tech_catalog_cache = json.load(f)
        else:
            _tech_catalog_cache = {"categories": []}
    return _tech_catalog_cache


def _to_dict(resume: Union[Resume, Dict[str, Any]]) -> Dict[str, Any]:
    if isinstance(resume, Resume):
        return resume.model_dump()
    elif isinstance(resume, dict):
        return resume
    else:
        raise ValueError(f"Expected Resume model or dict, got {type(resume).__name__}")


def group_skills(skills: List[Union[Dict[str, Any], Any]]) -> List[Dict[str, Any]]:
    """Group skills by category following the tech_catalog order."""
    catalog = _load_catalog()
    categories = catalog.get("categories", [])
    cat_order = [c["id"] for c in categories]
    cat_names = {c["id"]: c["name"] for c in categories}

    grouped: Dict[str, List[str]] = {cid: [] for cid in cat_order}
    other_skills: List[str] = []

    for item in skills:
        if isinstance(item, dict):
            name = item.get("name", "").strip()
            cid = item.get("category_id", "").strip().lower()
        else:
            name = str(item).strip()
            cid = ""

        if not name:
            continue

        if cid in grouped:
            grouped[cid].append(name)
        else:
            other_skills.append(name)

    result: List[Dict[str, Any]] = []
    for cid in cat_order:
        if grouped[cid]:
            result.append({
                "category_id": cid,
                "category_name": cat_names.get(cid, cid.replace("_", " ").title()),
                "skills": grouped[cid],
            })

    if other_skills:
        result.append({
            "category_id": "other",
            "category_name": "Other Skills",
            "skills": other_skills,
        })

    return result


def prepare_links(links: List[Dict[str, Any]]) -> List[Dict[str, str]]:
    """Enrich social links with full URLs and user-facing labels."""
    prepared: List[Dict[str, str]] = []
    platforms = load_platforms()

    for item in links:
        pid = item.get("platform_id", "").strip().lower()
        val = item.get("value", "").strip()
        if not pid or not val:
            continue

        label = platforms.get(pid, {}).get("label", pid.capitalize())
        try:
            url = build_url(pid, val)
        except Exception:
            url = val

        prepared.append({
            "platform_id": pid,
            "label": label,
            "value": val,
            "url": url,
        })
    return prepared


def render_html(
    resume: Union[Resume, Dict[str, Any]],
    template_name: str = "classic",
    page_size: str = "A4",
) -> str:
    """Render resume to HTML string using the specified template."""
    data = _to_dict(resume)
    env = _get_jinja_env()

    clean_template = template_name.lower().strip()
    if clean_template not in ("classic", "compact", "technical"):
        clean_template = "classic"

    template_file = f"{clean_template}.html"
    tmpl = env.get_template(template_file)

    grouped = group_skills(data.get("skills", []))
    links = prepare_links(data.get("basics", {}).get("links", []))

    clean_page_size = "Letter" if page_size.lower() in ("letter", "us-letter") else "A4"

    return tmpl.render(
        resume=data,
        grouped_skills=grouped,
        links=links,
        page_size=clean_page_size,
    )


def _render_pdf_headless(html_content: str, page_size: str = "A4") -> bytes:
    """Render PDF using headless Edge or Chrome on host system."""
    # Find browser
    browser_candidates = [
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
        shutil.which("msedge") or "",
        shutil.which("chrome") or "",
        shutil.which("google-chrome") or "",
        shutil.which("chromium") or "",
    ]
    browser_path = next((p for p in browser_candidates if p and os.path.exists(p)), None)
    if not browser_path:
        raise RuntimeError("No headless browser found for PDF rendering fallback.")

    temp_dir = tempfile.mkdtemp(prefix="skillgap_pdf_")
    html_file = os.path.join(temp_dir, "resume.html")
    pdf_file = os.path.join(temp_dir, "resume.pdf")

    try:
        with open(html_file, "w", encoding="utf-8") as f:
            f.write(html_content)

        cmd = [
            browser_path,
            "--headless=new",
            "--disable-gpu",
            "--no-pdf-header-footer",
            f"--print-to-pdf={pdf_file}",
            html_file,
        ]
        result = subprocess.run(cmd, capture_output=True, timeout=30)
        if result.returncode != 0 or not os.path.exists(pdf_file):
            raise RuntimeError(f"Browser PDF generation failed: {result.stderr.decode('utf-8', errors='ignore')}")

        with open(pdf_file, "rb") as f:
            pdf_bytes = f.read()

        return pdf_bytes
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


def render_pdf(
    resume: Union[Resume, Dict[str, Any]],
    template_name: str = "classic",
    page_size: str = "A4",
) -> bytes:
    """Render resume to PDF bytes using WeasyPrint with headless browser fallback."""
    html_content = render_html(resume, template_name=template_name, page_size=page_size)

    # 1. Try native WeasyPrint
    try:
        import weasyprint
        return weasyprint.HTML(string=html_content).write_pdf()
    except (ImportError, OSError) as exc:
        logger.info(f"WeasyPrint native rendering unavailable ({exc}), using headless browser fallback.")
        return _render_pdf_headless(html_content, page_size=page_size)


def render_docx(resume: Union[Resume, Dict[str, Any]]) -> bytes:
    """Render resume to ATS-safe DOCX bytes using python-docx."""
    data = _to_dict(resume)
    doc = docx.Document()

    # Configure standard page setup (single column, 0.6 inch margins)
    for sec in doc.sections:
        sec.top_margin = Inches(0.6)
        sec.bottom_margin = Inches(0.6)
        sec.left_margin = Inches(0.6)
        sec.right_margin = Inches(0.6)

    basics = data.get("basics", {})
    name = (basics.get("name") or "Candidate Name").strip()
    headline = (basics.get("headline") or "").strip()

    # Name as Heading 1
    h1 = doc.add_heading(name, level=1)
    h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if h1.runs:
        h1.runs[0].font.size = Pt(20)
        h1.runs[0].font.color.rgb = RGBColor(0x11, 0x11, 0x11)
        h1.runs[0].font.bold = True

    # Headline
    if headline:
        p_head = doc.add_paragraph()
        p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p_head.add_run(headline)
        run.italic = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    # Contact line
    contact_parts: List[str] = []
    if basics.get("email"):
        contact_parts.append(basics["email"])
    if basics.get("phone"):
        contact_parts.append(basics["phone"])
    loc = basics.get("location") or {}
    loc_str = ", ".join(filter(None, [loc.get("city"), loc.get("region_country")]))
    if loc_str:
        contact_parts.append(loc_str)
    for link in prepare_links(basics.get("links", [])):
        contact_parts.append(f"{link['label']}: {link['value']}")

    if contact_parts:
        p_contact = doc.add_paragraph()
        p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p_contact.add_run("  •  ".join(contact_parts))
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Helper for adding section headings
    def _add_section_heading(title: str):
        h = doc.add_heading(title, level=2)
        if h.runs:
            h.runs[0].font.size = Pt(13)
            h.runs[0].font.color.rgb = RGBColor(0x22, 0x22, 0x22)
            h.runs[0].font.bold = True

    # Summary
    summary = data.get("summary") or basics.get("summary")
    if summary:
        _add_section_heading("Summary")
        p_sum = doc.add_paragraph(summary)
        p_sum.paragraph_format.space_after = Pt(6)

    # Experience
    experience = data.get("experience", [])
    if experience:
        _add_section_heading("Experience")
        for exp in experience:
            p_exp = doc.add_paragraph()
            r_title = p_exp.add_run(exp.get("title", ""))
            r_title.bold = True
            r_co = p_exp.add_run(f" — {exp.get('company', '')}")
            r_co.italic = True
            if exp.get("location"):
                p_exp.add_run(f" ({exp['location']})")

            start = exp.get("start") or ""
            end = "Present" if exp.get("current") else (exp.get("end") or "Present")
            if start or end:
                r_dates = p_exp.add_run(f"  [{start} – {end}]")
                r_dates.font.size = Pt(9)
                r_dates.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

            p_exp.paragraph_format.space_after = Pt(2)
            for bullet in exp.get("bullets", []):
                doc.add_paragraph(bullet, style="List Bullet")

    # Education
    education = data.get("education", [])
    if education:
        _add_section_heading("Education")
        for edu in education:
            p_edu = doc.add_paragraph()
            r_sch = p_edu.add_run(edu.get("school", ""))
            r_sch.bold = True

            deg_parts = filter(None, [edu.get("degree"), edu.get("field")])
            deg_str = ", ".join(deg_parts)
            if deg_str:
                p_edu.add_run(f" — {deg_str}")
            if edu.get("score"):
                p_edu.add_run(f" (GPA/Score: {edu['score']})")

            start = edu.get("start") or ""
            end = edu.get("end") or ""
            if start or end:
                dates = f"  [{start} – {end}]" if (start and end) else f"  [{start or end}]"
                r_dates = p_edu.add_run(dates)
                r_dates.font.size = Pt(9)
                r_dates.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

            p_edu.paragraph_format.space_after = Pt(2)
            for detail in edu.get("details", []):
                doc.add_paragraph(detail, style="List Bullet")

    # Projects
    projects = data.get("projects", [])
    if projects:
        _add_section_heading("Projects")
        for proj in projects:
            p_proj = doc.add_paragraph()
            r_name = p_proj.add_run(proj.get("name", ""))
            r_name.bold = True
            if proj.get("link"):
                p_proj.add_run(f" ({proj['link']})")
            if proj.get("tech"):
                r_tech = p_proj.add_run(f" — Stack: {', '.join(proj['tech'])}")
                r_tech.italic = True

            p_proj.paragraph_format.space_after = Pt(2)
            for bullet in proj.get("bullets", []):
                doc.add_paragraph(bullet, style="List Bullet")

    # Skills
    grouped = group_skills(data.get("skills", []))
    if grouped:
        _add_section_heading("Skills")
        for g in grouped:
            p_sk = doc.add_paragraph()
            r_lbl = p_sk.add_run(f"{g['category_name']}: ")
            r_lbl.bold = True
            p_sk.add_run(", ".join(g["skills"]))
            p_sk.paragraph_format.space_after = Pt(2)

    # Certifications
    certifications = data.get("certifications", [])
    if certifications:
        _add_section_heading("Certifications")
        for cert in certifications:
            p_cert = doc.add_paragraph(style="List Bullet")
            r_name = p_cert.add_run(cert.get("name", ""))
            r_name.bold = True
            if cert.get("issuer"):
                p_cert.add_run(f" — {cert['issuer']}")
            if cert.get("date"):
                p_cert.add_run(f" ({cert['date']})")

    # Achievements
    achievements = data.get("achievements", [])
    if achievements:
        _add_section_heading("Achievements")
        for ach in achievements:
            doc.add_paragraph(ach, style="List Bullet")

    # Publications
    publications = data.get("publications", [])
    if publications:
        _add_section_heading("Publications")
        for pub in publications:
            p_pub = doc.add_paragraph(style="List Bullet")
            r_t = p_pub.add_run(pub.get("title", ""))
            r_t.bold = True
            if pub.get("venue"):
                p_pub.add_run(f" — {pub['venue']}")
            if pub.get("year"):
                p_pub.add_run(f" ({pub['year']})")
            if pub.get("link"):
                p_pub.add_run(f" [{pub['link']}]")

    # Custom sections
    custom_sections = data.get("custom_sections", [])
    for cs in custom_sections:
        _add_section_heading(cs.get("title", "Custom Section"))
        for b in cs.get("bullets", []):
            doc.add_paragraph(b, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_markdown(resume: Union[Resume, Dict[str, Any]]) -> str:
    """Render resume to clean Markdown text."""
    data = _to_dict(resume)
    lines: List[str] = []

    basics = data.get("basics", {})
    name = (basics.get("name") or "Candidate Name").strip()
    lines.append(f"# {name}\n")

    if basics.get("headline"):
        lines.append(f"**{basics['headline']}**\n")

    contact_parts: List[str] = []
    if basics.get("email"):
        contact_parts.append(f"[{basics['email']}](mailto:{basics['email']})")
    if basics.get("phone"):
        contact_parts.append(basics["phone"])
    loc = basics.get("location") or {}
    loc_str = ", ".join(filter(None, [loc.get("city"), loc.get("region_country")]))
    if loc_str:
        contact_parts.append(loc_str)
    for link in prepare_links(basics.get("links", [])):
        contact_parts.append(f"[{link['label']}]({link['url']})")

    if contact_parts:
        lines.append(" | ".join(contact_parts) + "\n")

    summary = data.get("summary") or basics.get("summary")
    if summary:
        lines.append(f"## Summary\n\n{summary}\n")

    experience = data.get("experience", [])
    if experience:
        lines.append("## Experience\n")
        for exp in experience:
            start = exp.get("start") or ""
            end = "Present" if exp.get("current") else (exp.get("end") or "Present")
            dates = f"({start} – {end})" if (start or end) else ""
            lines.append(f"### {exp.get('title')} — {exp.get('company')} {dates}\n")
            for b in exp.get("bullets", []):
                lines.append(f"- {b}")
            lines.append("")

    education = data.get("education", [])
    if education:
        lines.append("## Education\n")
        for edu in education:
            deg = ", ".join(filter(None, [edu.get("degree"), edu.get("field")]))
            lines.append(f"### {edu.get('school')} — {deg}\n")
            for d in edu.get("details", []):
                lines.append(f"- {d}")
            lines.append("")

    projects = data.get("projects", [])
    if projects:
        lines.append("## Projects\n")
        for proj in projects:
            tech = f" (Tech: {', '.join(proj['tech'])})" if proj.get("tech") else ""
            lines.append(f"### {proj.get('name')}{tech}\n")
            for b in proj.get("bullets", []):
                lines.append(f"- {b}")
            lines.append("")

    grouped = group_skills(data.get("skills", []))
    if grouped:
        lines.append("## Skills\n")
        for g in grouped:
            lines.append(f"- **{g['category_name']}:** {', '.join(g['skills'])}")
        lines.append("")

    certifications = data.get("certifications", [])
    if certifications:
        lines.append("## Certifications\n")
        for cert in certifications:
            issuer = f" — {cert['issuer']}" if cert.get("issuer") else ""
            date = f" ({cert['date']})" if cert.get("date") else ""
            lines.append(f"- {cert.get('name')}{issuer}{date}")
        lines.append("")

    achievements = data.get("achievements", [])
    if achievements:
        lines.append("## Achievements\n")
        for ach in achievements:
            lines.append(f"- {ach}")
        lines.append("")

    publications = data.get("publications", [])
    if publications:
        lines.append("## Publications\n")
        for pub in publications:
            venue = f" — {pub['venue']}" if pub.get("venue") else ""
            year = f" ({pub['year']})" if pub.get("year") else ""
            lines.append(f"- {pub.get('title')}{venue}{year}")
        lines.append("")

    custom_sections = data.get("custom_sections", [])
    for cs in custom_sections:
        lines.append(f"## {cs.get('title', 'Custom Section')}\n")
        for b in cs.get("bullets", []):
            lines.append(f"- {b}")
        lines.append("")

    return "\n".join(lines).strip() + "\n"


def render_text(resume: Union[Resume, Dict[str, Any]]) -> str:
    """Render resume to clean plaintext with ASCII bullet points."""
    data = _to_dict(resume)
    lines: List[str] = []

    basics = data.get("basics", {})
    name = (basics.get("name") or "Candidate Name").strip()
    lines.append(name.upper())

    if basics.get("headline"):
        lines.append(basics["headline"])

    contact_parts: List[str] = []
    if basics.get("email"):
        contact_parts.append(basics["email"])
    if basics.get("phone"):
        contact_parts.append(basics["phone"])
    loc = basics.get("location") or {}
    loc_str = ", ".join(filter(None, [loc.get("city"), loc.get("region_country")]))
    if loc_str:
        contact_parts.append(loc_str)
    for link in prepare_links(basics.get("links", [])):
        contact_parts.append(f"{link['label']}: {link['value']}")

    if contact_parts:
        lines.append(" | ".join(contact_parts))

    lines.append("-" * 60)

    summary = data.get("summary") or basics.get("summary")
    if summary:
        lines.append("SUMMARY")
        lines.append(summary)
        lines.append("")

    experience = data.get("experience", [])
    if experience:
        lines.append("EXPERIENCE")
        for exp in experience:
            start = exp.get("start") or ""
            end = "Present" if exp.get("current") else (exp.get("end") or "Present")
            dates = f" [{start} - {end}]" if (start or end) else ""
            lines.append(f"{exp.get('title')} - {exp.get('company')}{dates}")
            for b in exp.get("bullets", []):
                lines.append(f"  * {b}")
            lines.append("")

    education = data.get("education", [])
    if education:
        lines.append("EDUCATION")
        for edu in education:
            deg = ", ".join(filter(None, [edu.get("degree"), edu.get("field")]))
            lines.append(f"{edu.get('school')} - {deg}")
            for d in edu.get("details", []):
                lines.append(f"  * {d}")
            lines.append("")

    projects = data.get("projects", [])
    if projects:
        lines.append("PROJECTS")
        for proj in projects:
            tech = f" (Tech: {', '.join(proj['tech'])})" if proj.get("tech") else ""
            lines.append(f"{proj.get('name')}{tech}")
            for b in proj.get("bullets", []):
                lines.append(f"  * {b}")
            lines.append("")

    grouped = group_skills(data.get("skills", []))
    if grouped:
        lines.append("SKILLS")
        for g in grouped:
            lines.append(f"  * {g['category_name']}: {', '.join(g['skills'])}")
        lines.append("")

    certifications = data.get("certifications", [])
    if certifications:
        lines.append("CERTIFICATIONS")
        for cert in certifications:
            issuer = f" - {cert['issuer']}" if cert.get("issuer") else ""
            date = f" ({cert['date']})" if cert.get("date") else ""
            lines.append(f"  * {cert.get('name')}{issuer}{date}")
        lines.append("")

    achievements = data.get("achievements", [])
    if achievements:
        lines.append("ACHIEVEMENTS")
        for ach in achievements:
            lines.append(f"  * {ach}")
        lines.append("")

    publications = data.get("publications", [])
    if publications:
        lines.append("PUBLICATIONS")
        for pub in publications:
            lines.append(f"  * {pub.get('title')}")
        lines.append("")

    custom_sections = data.get("custom_sections", [])
    for cs in custom_sections:
        lines.append(cs.get("title", "CUSTOM SECTION").upper())
        for b in cs.get("bullets", []):
            lines.append(f"  * {b}")
        lines.append("")

    return "\n".join(lines).strip() + "\n"


def export_json(resume: Union[Resume, Dict[str, Any]]) -> str:
    """Export schema-compliant JSON string."""
    data = _to_dict(resume)
    # Validate with Resume model to ensure strict schema compliance
    validated = Resume.model_validate(data)
    return validated.model_dump_json(indent=2)
